import os
import re
import logging
from pathlib import Path
from openpyxl import load_workbook
from icrawler.builtin import BingImageCrawler

# Silence verbose network logs so your terminal remains clean
logging.getLogger('icrawler').setLevel(logging.WARNING)

# Pipeline Configuration
EXCEL_FILE = str(Path(__file__).resolve().parent / "emadini" / "e-waste datasets.xlsx")
OUTPUT_DIR = "ewaste_visual_dataset"
IMAGES_PER_DEVICE = 30

os.makedirs(OUTPUT_DIR, exist_ok=True)

def sanitize_path(text):
    """Filters out illegal operating system characters to ensure folders save properly."""
    return re.sub(r'[\\/*?:"<>| ]', '_', str(text).strip())

def main():
    global EXCEL_FILE
    # Check if the Excel file exists; if not, try to find one in current directory
    if not os.path.exists(EXCEL_FILE):
        xlsx_files = [f for f in os.listdir('.') if f.endswith('.xlsx')]
        if xlsx_files:
            EXCEL_FILE = xlsx_files[0]
            print(f"📂 Found Excel file automatically: using '{EXCEL_FILE}'")
        else:
            print(f"❌ Error: Could not find an Excel file in this folder.")
            return

    print(f"📂 Opening workbook: {EXCEL_FILE}...")
    wb = load_workbook(filename=EXCEL_FILE, read_only=True, data_only=True)
    sheet = wb.active
    
    header_row = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    
    try:
        item_col_idx = header_row.index("Item Name") + 1
        brand_col_idx = header_row.index("Brand Name") + 1
        print(f"✅ Columns mapped: 'Item Name' (Col {item_col_idx}) & 'Brand Name' (Col {brand_col_idx}).")
    except ValueError:
        print("⚠️ Warning: Exact columns 'Item Name' or 'Brand Name' not found in row 1. Using default columns 1 and 3.")
        item_col_idx = 1
        brand_col_idx = 3

    row_count = 0
    for row in sheet.iter_rows(min_row=2):
        row_count += 1
        # Extract values, handle potential None
        row_values = [cell.value for cell in row]
        item_name = row_values[item_col_idx - 1] if len(row_values) >= item_col_idx else None
        brand_name = row_values[brand_col_idx - 1] if len(row_values) >= brand_col_idx else None

        # Skip if essential data missing
        if not item_name:
            continue

        # Compose search query
        search_query = f"{item_name} {brand_name}" if brand_name else str(item_name)

        # Check if images already exist for this item
        target_folder = os.path.join(OUTPUT_DIR, sanitize_path(search_query))
        if os.path.exists(target_folder):
            if row_count % 100 == 0:
                print(f"🔄 Checkpoint: Folder for '{search_query}' already exists. Skipping download.")
            continue
        else:
            os.makedirs(target_folder, exist_ok=True)

        print(f"🚀 Downloading {IMAGES_PER_DEVICE} images for row [{row_count}]: '{search_query}'")
        try:
            crawler = BingImageCrawler(downloader_threads=4, storage={'root_dir': target_folder})
            crawler.crawl(keyword=search_query, max_num=IMAGES_PER_DEVICE)
        except Exception as e:
            print(f"⚠️ Row {row_count} skipped due to error: {e}")
            continue

    print("\n🏁 Process complete! Your dataset has been generated.")

if __name__ == "__main__":
    main()